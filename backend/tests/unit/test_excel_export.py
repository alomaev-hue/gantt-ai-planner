from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.domain.operations import apply_operations, operations_adapter
from app.domain.scheduler import schedule
from app.domain.seed import build_demo_plan
from app.excel.export import EXPORT_HEADERS, export_plan_xlsx
from app.excel.parse import parse_plan_xlsx

TODAY = date(2026, 9, 25)


def test_export_layout():
    sp = schedule(build_demo_plan(TODAY))
    wb = load_workbook(BytesIO(export_plan_xlsx(sp)))
    ws = wb["План"]
    assert [c.value for c in ws[1]] == EXPORT_HEADERS
    assert ws.freeze_panes == "A2"
    assert ws.max_row == len(sp.tasks) + 1
    first = sp.tasks[0]
    row = [c.value for c in ws[2]]
    assert row[0] == first.id and row[1] == first.name
    assert row[7].date() == first.start and row[8].date() == first.end


def test_round_trip_preserves_plan_including_constraints_and_lags():
    plan = build_demo_plan(TODAY)
    res = apply_operations(
        plan,
        operations_adapter.validate_python(
            [
                {"op": "move_task", "id": 21, "start_date": "2026-10-12"},
                {"op": "delete_task", "id": 7},
            ]
        ),
    )
    sp = res.scheduled
    imported = parse_plan_xlsx(export_plan_xlsx(sp), sp.project_start)
    assert imported.ok, imported.errors
    p2 = imported.plan
    assert [t.model_dump() for t in p2.tasks] == [t.model_dump() for t in res.plan.tasks]
    key = lambda d: (d.predecessor_id, d.successor_id, d.lag)  # noqa: E731
    assert sorted(map(key, p2.dependencies)) == sorted(map(key, res.plan.dependencies))
    sp2 = schedule(p2)
    assert [(t.start, t.end) for t in sp2.tasks] == [(t.start, t.end) for t in sp.tasks]


def test_formula_like_text_is_exported_as_plain_string_and_round_trips():
    plan = build_demo_plan(TODAY)
    res = apply_operations(
        plan,
        operations_adapter.validate_python(
            [
                {
                    "op": "update_task",
                    "id": 1,
                    "name": '=HYPERLINK("http://evil")',
                    "description": "=1+1",
                    "assignee": "=cmd|' /C calc'!A0",
                }
            ]
        ),
    )
    data = export_plan_xlsx(res.scheduled)
    ws = load_workbook(BytesIO(data))["План"]
    for cell in ws[2][1:4]:
        assert cell.data_type == "s", cell.value
        assert cell.value.startswith("=")
    imported = parse_plan_xlsx(data, res.scheduled.project_start)
    assert imported.ok, imported.errors
    task = imported.plan.tasks[0]
    assert task.name == '=HYPERLINK("http://evil")'
    assert task.description == "=1+1"
    assert task.assignee == "=cmd|' /C calc'!A0"
