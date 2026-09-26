from datetime import date
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


GOOD = xlsx(
    [
        ["Задача", "Описание", "Исполнитель", "Длительность", "Предшественники"],
        ["A", "", "Олег", 2, None],
        ["B", "", "Олег", 3, "1"],
    ]
)


async def test_import_replaces_plan_and_is_undoable(session_client):
    r = await session_client.post(
        "/api/plan/import",
        files={"file": ("office.xlsx", GOOD, XLSX_MIME)},
        data={"project_start": "2026-10-04"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["ok"] is True and len(body["plan"]["plan"]["tasks"]) == 2
    assert body["plan"]["plan"]["project_start"] == "2026-10-05"
    history = (await session_client.post("/api/plan/undo")).json()
    assert len(history["plan"]["tasks"]) == 25


async def test_import_errors_and_size_limit(session_client):
    bad = xlsx([["Задача", "Длительность", "Предшественники"], ["A", 1, "7"]])
    r = await session_client.post(
        "/api/plan/import",
        files={"file": ("b.xlsx", bad, XLSX_MIME)},
        data={"project_start": "2026-10-05"},
    )
    assert r.status_code == 422 and r.json()["ok"] is False and r.json()["errors"][0]["row"] == 2
    huge = b"0" * (2 * 1024 * 1024 + 1)
    r = await session_client.post(
        "/api/plan/import",
        files={"file": ("h.xlsx", huge, XLSX_MIME)},
        data={"project_start": "2026-10-05"},
    )
    assert r.status_code == 413 and r.json()["error"]["code"] == "file_too_large"


async def test_export_filename_uses_the_clients_date(session_client):
    r = await session_client.get("/api/plan/export", params={"today": "2026-12-31"})
    assert 'filename="plan-2026-12-31.xlsx"' in r.headers["content-disposition"]


@pytest.mark.parametrize("bad", ["2026-13-01", "31.12.2026", 'x"; evil=1', ""])
async def test_export_filename_falls_back_to_server_date_on_bad_input(session_client, bad):
    r = await session_client.get("/api/plan/export", params={"today": bad})
    assert r.status_code == 200
    assert f'filename="plan-{date.today():%Y-%m-%d}.xlsx"' in r.headers["content-disposition"]


async def test_export_downloads_xlsx(session_client):
    r = await session_client.get("/api/plan/export")
    assert r.status_code == 200 and r.headers["content-type"].startswith(XLSX_MIME)
    assert 'filename="plan-' in r.headers["content-disposition"]
    assert load_workbook(BytesIO(r.content))["План"].max_row == 26
