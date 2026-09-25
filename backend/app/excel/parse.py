"""Tolerant .xlsx plan importer (spec §9.1). Collects all issues instead of failing fast."""

import math
import re
import zipfile
from datetime import date, datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import BaseModel, ValidationError

from app.domain.calendar import next_workday
from app.domain.errors import CycleError
from app.domain.models import MAX_TASKS, Dependency, Plan, Task
from app.domain.scheduler import topological_order
from app.excel.headers import match_column

LIMITS = {"name": 200, "description": 2000, "assignee": 100}
_DURATION_RE = re.compile(
    r"^(\d+(?:[.,]\d+)?)\s*(д|дн|дня|дней|день|d|day|days|н|нед|недел[яьи]|w|wk|week|weeks)?\.?$"
)
_NUM_REF_RE = re.compile(r"^(\d+)(?:[.,]0+)?\s*(?:fs|он)?\s*(?:\+\s*(\d+)\s*(?:д|дн|d)?\.?)?$")
_NAME_LAG_RE = re.compile(r"^(.*?)\s*\+\s*(\d+)\s*(?:д|дн|d)?\.?$")
_OTHER_TYPES_RE = re.compile(r"^\d+\s*(ss|ff|sf|нн|оо|но)\b")


class ImportIssue(BaseModel):
    row: int | None
    message: str


class ImportResult(BaseModel):
    ok: bool
    plan: Plan | None
    errors: list[ImportIssue]
    warnings: list[ImportIssue]


def parse_duration(value: object) -> tuple[int, str | None]:
    if value is None or (isinstance(value, str) and not value.strip()):
        return 1, "длительность не указана, принята 1 день"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        number, unit = float(value), None
    else:
        m = _DURATION_RE.match(str(value).strip().casefold())
        if not m:
            raise ValueError(f"не удалось распознать длительность «{value}»")
        number, unit = float(m.group(1).replace(",", ".")), m.group(2)
    if unit and unit[0] in ("н", "w"):
        number *= 5
    if number <= 0:
        return 1, "длительность 0, принята 1 день"
    days = math.ceil(number)
    warning = None if days == number else f"длительность {number:g} округлена до {days}"
    if days > 999:
        raise ValueError("длительность больше 999 дней")
    return days, warning


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _text(value)
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _fail(errors: list[ImportIssue], warnings: list[ImportIssue]) -> ImportResult:
    return ImportResult(ok=False, plan=None, errors=errors, warnings=warnings)


def _short(exc: ValidationError) -> str:
    return "; ".join(f"{'.'.join(map(str, e['loc']))}: {e['msg']}" for e in exc.errors())


def parse_plan_xlsx(data: bytes, project_start: date) -> ImportResult:
    errors: list[ImportIssue] = []
    warnings: list[ImportIssue] = []
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError, ValueError):
        return _fail([ImportIssue(row=None, message="Файл не является корректным .xlsx")], [])
    ws = wb.worksheets[0]
    rows = [(i, list(r)) for i, r in enumerate(ws.iter_rows(values_only=True), start=1)]
    wb.close()

    header_row: int | None = None
    columns: dict[str, int] = {}
    for row_no, values in rows[:30]:
        mapped = {key: idx for idx, v in enumerate(values) if (key := match_column(v))}
        if "name" in mapped:
            header_row, columns = row_no, mapped
            break
    if header_row is None or "duration" not in columns:
        missing = "«Задача»" if header_row is None else "«Длительность»"
        return _fail(
            [ImportIssue(row=None, message=f"Не найдена обязательная колонка {missing}")], []
        )

    def cell(values: list[Any], key: str) -> Any:
        idx = columns.get(key)
        return values[idx] if idx is not None and idx < len(values) else None

    raw_tasks: list[dict[str, Any]] = []
    for row_no, values in rows:
        if row_no <= header_row or all(_text(v) == "" for v in values):
            continue
        name = _text(cell(values, "name"))
        if not name:
            errors.append(ImportIssue(row=row_no, message="Пустое название задачи"))
            continue
        item: dict[str, Any] = {"row": row_no, "name": name}
        for key in ("name", "description", "assignee"):
            text = name if key == "name" else _text(cell(values, key))
            if len(text) > LIMITS[key]:
                text = text[: LIMITS[key]]
                warnings.append(
                    ImportIssue(
                        row=row_no, message=f"Поле «{key}» обрезано до {LIMITS[key]} символов"
                    )
                )
            item[key] = text
        try:
            item["duration"], warn = parse_duration(cell(values, "duration"))
            if warn:
                warnings.append(ImportIssue(row=row_no, message=warn.capitalize()))
        except ValueError as exc:
            errors.append(ImportIssue(row=row_no, message=str(exc).capitalize()))
            continue
        item["number"] = cell(values, "number")
        item["preds"] = _text(cell(values, "predecessors"))
        raw_constraint = cell(values, "constraint")
        item["constraint"] = _as_date(raw_constraint) if _text(raw_constraint) else None
        if _text(raw_constraint) and item["constraint"] is None:
            warnings.append(
                ImportIssue(
                    row=row_no, message=f"Дата «{raw_constraint}» не распознана и пропущена"
                )
            )
        raw_tasks.append(item)

    if not raw_tasks and not errors:
        return _fail([ImportIssue(row=None, message="В файле нет задач")], warnings)
    if len(raw_tasks) > MAX_TASKS:
        return _fail(
            [*errors, ImportIssue(row=None, message=f"Больше {MAX_TASKS} задач в файле")],
            warnings,
        )

    # ids: from «№» column if present, else 1..n
    use_numbers = "number" in columns
    ids: dict[int, int] = {}  # row -> id
    seen: dict[int, int] = {}
    for position, item in enumerate(raw_tasks, start=1):
        if not use_numbers:
            ids[item["row"]] = position
            continue
        text = _text(item["number"])
        if not text.isdigit() or int(text) < 1:
            errors.append(
                ImportIssue(row=item["row"], message="«№» должен быть положительным целым числом")
            )
            continue
        number = int(text)
        if number in seen:
            errors.append(
                ImportIssue(
                    row=item["row"], message=f"№ {number} повторяется (строка {seen[number]})"
                )
            )
            continue
        seen[number] = item["row"]
        ids[item["row"]] = number

    by_position = {pos: ids.get(it["row"]) for pos, it in enumerate(raw_tasks, start=1)}
    by_name: dict[str, list[int]] = {}
    for it in raw_tasks:
        if it["row"] in ids:
            by_name.setdefault(it["name"].casefold(), []).append(ids[it["row"]])
    known_ids = set(ids.values())

    deps: dict[tuple[int, int], int] = {}
    for it in raw_tasks:
        row, succ = it["row"], ids.get(it["row"])
        if succ is None or not it["preds"]:
            continue
        for token in (t.strip() for t in re.split(r"[;,\n]", it["preds"]) if t.strip()):
            ref, lag = _resolve(token, use_numbers, by_position, by_name, known_ids)
            if isinstance(ref, str):
                errors.append(ImportIssue(row=row, message=ref))
                continue
            if ref == succ:
                errors.append(
                    ImportIssue(row=row, message="Задача не может зависеть от самой себя")
                )
                continue
            if lag > 365:
                errors.append(ImportIssue(row=row, message=f"Лаг больше 365 дней в «{token}»"))
                continue
            deps[(ref, succ)] = max(lag, deps.get((ref, succ), 0))

    if errors:
        return _fail(errors, warnings)

    try:
        plan = Plan(
            project_start=next_workday(project_start),
            tasks=[
                Task(
                    id=ids[it["row"]],
                    name=it["name"],
                    description=it["description"],
                    assignee=it["assignee"] or None,
                    duration=it["duration"],
                    constraint_start=it["constraint"],
                )
                for it in raw_tasks
            ],
            dependencies=[
                Dependency(predecessor_id=p, successor_id=s, lag=lag)
                for (p, s), lag in deps.items()
            ],
        )
    except ValidationError as exc:
        return _fail(
            [ImportIssue(row=None, message=f"Некорректные данные: {_short(exc)}")], warnings
        )
    try:
        topological_order(plan)
    except CycleError as exc:
        rows_by_id = {v: k for k, v in ids.items()}
        cycle_rows = ", ".join(str(rows_by_id[i]) for i in exc.cycle)
        return _fail(
            [ImportIssue(row=None, message=f"{exc.message} (строки {cycle_rows})")], warnings
        )
    return ImportResult(ok=True, plan=plan, errors=[], warnings=warnings)


def _resolve(
    token: str,
    use_numbers: bool,
    by_position: dict[int, int | None],
    by_name: dict[str, list[int]],
    known_ids: set[int],
) -> tuple[int | str, int]:
    low = token.casefold()
    if _OTHER_TYPES_RE.match(low):
        return "Поддерживается только тип связи «окончание–начало»", 0
    m = _NUM_REF_RE.match(low)
    if m:
        number, lag = int(m.group(1)), int(m.group(2) or 0)
        ref = number if use_numbers else by_position.get(number)
        if ref is None or ref not in known_ids:
            return f"Предшественник «{token}» не найден", 0
        return ref, lag
    name, lag = token, 0
    m = _NAME_LAG_RE.match(token)
    if m:
        name, lag = m.group(1), int(m.group(2))
    matches = by_name.get(name.strip().casefold(), [])
    if not matches:
        return f"Предшественник «{token}» не найден", 0
    if len(matches) > 1:
        return f"Название «{name}» встречается несколько раз — укажите №", 0
    return matches[0], lag
