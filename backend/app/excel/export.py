from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.domain.diff import format_predecessors
from app.domain.scheduler import ScheduledPlan

EXPORT_HEADERS = [
    "№",
    "Задача",
    "Описание",
    "Исполнитель",
    "Длительность",
    "Предшественники",
    "Не раньше",
    "Начало",
    "Окончание",
    "Резерв",
    "Критическая",
]
_WIDTHS = [6, 40, 50, 22, 14, 18, 13, 13, 13, 9, 12]
_DATE_COLUMNS = (7, 8, 9)  # 1-based: «Не раньше», «Начало», «Окончание»


def export_plan_xlsx(sp: ScheduledPlan) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "План"
    ws.append(EXPORT_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for t in sp.tasks:
        ws.append(
            [
                t.id,
                t.name,
                t.description,
                t.assignee or "",
                t.duration,
                format_predecessors(sp.dependencies, t.id),
                t.constraint_start,
                t.start,
                t.end,
                t.slack,
                "да" if t.is_critical else "",
            ]
        )
    for row in ws.iter_rows(min_row=2):
        for col in _DATE_COLUMNS:
            row[col - 1].number_format = "DD.MM.YYYY"
    for i, width in enumerate(_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
