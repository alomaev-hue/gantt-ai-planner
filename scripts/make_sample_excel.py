"""Generate examples/sample-plan.xlsx — a fictional «Переезд офиса» plan.

Exercises the tolerant .xlsx importer (see backend/app/excel/parse.py):
  - a title row above the real header (header detection must skip it)
  - mixed duration formats: plain number, "<n>д", "<n> нед", "<n> дней"
  - predecessors referenced by row/task number, by task name, with a lag
    ("<n>+<lag>"), and multiple predecessors separated by ";"

Run from the backend project (needs the openpyxl dependency):
    cd backend && uv run python ../scripts/make_sample_excel.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUTPUT = Path(__file__).resolve().parent.parent / "examples" / "sample-plan.xlsx"

TITLE = "План переезда офиса"
HEADER = ["Задача", "Описание", "Исполнитель", "Длительность", "Предшественники"]

OLEG = "Олег Сидоров"
NATALIA = "Наталья Белова"
PAVEL = "Павел Громов"
IRINA = "Ирина Лебедева"

# (name, description, assignee, duration, predecessors)
TASKS: list[tuple[str, str, str, object, str]] = [
    ("Планирование переезда", "Составить график и бюджет переезда", OLEG, 3, ""),
    ("Упаковка мебели", "Упаковать мебель для транспортировки", NATALIA, "2д", "1"),
    ("Упаковка документов", "Упаковать архив и текущие документы", PAVEL, "1 нед", "1"),
    (
        "Демонтаж мебели",
        "Разобрать крупную мебель перед вывозом",
        IRINA,
        "5 дней",
        "Упаковка мебели",
    ),
    ("Заказ грузового транспорта", "Забронировать грузовой автомобиль и грузчиков", OLEG, 2, "1"),
    ("Подготовка нового офиса", "Проверить готовность нового помещения", NATALIA, 3, ""),
    ("Уборка нового офиса", "Провести генеральную уборку перед заездом", PAVEL, 2, "6"),
    (
        "Демонтаж IT-инфраструктуры",
        "Отключить и упаковать сервера и сетевое оборудование",
        IRINA,
        2,
        "3+1",
    ),
    ("Погрузка мебели", "Погрузить упакованную и демонтированную мебель", OLEG, 1, "4; 5"),
    ("Перевозка мебели", "Перевезти мебель в новый офис", NATALIA, 1, "9"),
    (
        "Разгрузка и расстановка мебели",
        "Разгрузить и расставить мебель по новому плану",
        PAVEL,
        2,
        "7; 10",
    ),
    (
        "Перевозка и установка IT-инфраструктуры",
        "Перевезти сервера и оборудование в новый офис",
        IRINA,
        2,
        "8",
    ),
    ("Настройка сети и оборудования", "Настроить сеть, серверы и рабочие места", OLEG, "2д", "12"),
    ("Переезд сотрудников", "Переселить сотрудников на новые рабочие места", NATALIA, 1, "11; 13"),
    ("Итоговая проверка и приёмка офиса", "Проверить готовность офиса к работе", PAVEL, 1, "14"),
]


def build() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "План"

    ws.append([TITLE])
    ws["A1"].font = Font(bold=True, size=14)

    ws.append(HEADER)
    for cell in ws[2]:
        cell.font = Font(bold=True)

    for row in TASKS:
        ws.append(list(row))

    widths = [36, 44, 20, 14, 18]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=2, column=idx).column_letter].width = width

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(f"wrote {OUTPUT} ({len(TASKS)} tasks)")


if __name__ == "__main__":
    build()
